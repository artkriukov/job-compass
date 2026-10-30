from io import BytesIO
from pathlib import Path

import pytest
from docx import Document
from fastapi.testclient import TestClient

from app.main import app, get_resume_store
from app.services.resume_store import ResumeStore, ResumeValidationError, extract_resume_text

RESUME_TEXT = """Артём Крюков

Опыт работы

Разработчик программного обеспечения с опытом создания веб-приложений и API.

Навыки

Python, TypeScript, React, PostgreSQL, Docker, тестирование и автоматизация.
"""


def test_resume_store_saves_text_and_metadata(tmp_path: Path) -> None:
    store = ResumeStore(tmp_path)

    snapshot = store.save(
        filename="resume.txt",
        content_type="text/plain",
        payload=RESUME_TEXT.encode(),
    )

    assert snapshot.uploaded is True
    assert snapshot.original_filename == "resume.txt"
    assert snapshot.extracted_characters == len(RESUME_TEXT.strip())
    assert store.load().preview == RESUME_TEXT.strip()
    assert len(list((tmp_path / "documents").glob("resume-*.txt"))) == 2


def test_resume_extractor_rejects_unsupported_format() -> None:
    with pytest.raises(ResumeValidationError, match="PDF, DOCX, TXT и MD"):
        extract_resume_text("resume.rtf", RESUME_TEXT.encode())


def test_resume_extractor_reads_docx() -> None:
    document = Document()
    document.add_heading("Артём Крюков", level=1)
    document.add_heading("Опыт работы", level=2)
    document.add_paragraph(RESUME_TEXT)
    payload = BytesIO()
    document.save(payload)

    extracted = extract_resume_text("resume.docx", payload.getvalue())

    assert extracted.startswith("Артём Крюков\n\nОпыт работы")
    assert "Python, TypeScript" in extracted


def test_resume_endpoints_upload_and_read(tmp_path: Path) -> None:
    store = ResumeStore(tmp_path)
    app.dependency_overrides[get_resume_store] = lambda: store
    client = TestClient(app)

    empty = client.get("/api/career-ops/resume")
    uploaded = client.post(
        "/api/career-ops/resume",
        files={"file": ("resume.md", RESUME_TEXT.encode(), "text/markdown")},
    )
    loaded = client.get("/api/career-ops/resume")

    app.dependency_overrides.clear()
    assert empty.json() == {
        "uploaded": False,
        "original_filename": None,
        "content_type": None,
        "size_bytes": None,
        "extracted_characters": None,
        "preview": None,
        "uploaded_at": None,
    }
    assert uploaded.status_code == 200
    assert uploaded.json()["original_filename"] == "resume.md"
    assert loaded.json()["preview"].startswith("Артём Крюков")


def test_resume_endpoint_rejects_unsupported_format(tmp_path: Path) -> None:
    app.dependency_overrides[get_resume_store] = lambda: ResumeStore(tmp_path)
    client = TestClient(app)

    response = client.post(
        "/api/career-ops/resume",
        files={"file": ("resume.rtf", RESUME_TEXT.encode(), "application/rtf")},
    )

    app.dependency_overrides.clear()
    assert response.status_code == 422
    assert "PDF, DOCX, TXT и MD" in response.json()["detail"]
